import threading
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from agent_ai_builder import datos
from agent_ai_builder.tools import _PRODUCTS


NUM_CLIENTES = 10
codigos: list[str] = []
errores: list[str] = []
candado_resultados = threading.Lock()


def registrar_pedido_de_cliente(numero: int) -> None:
    try:
        respuesta = datos.registrar_pedido(
            [{"producto": "espresso", "cantidad": 1}],
            _PRODUCTS,
            contacto=f"cliente-{numero}",
        )
        if not respuesta.startswith("Pedido ") or "registrado" not in respuesta:
            raise RuntimeError(respuesta)

        codigo = respuesta.split()[1]
        with candado_resultados:
            codigos.append(codigo)
    except Exception as error:
        with candado_resultados:
            errores.append(f"Cliente {numero}: {error}")


with TemporaryDirectory() as carpeta_temporal:
    RUTA_PEDIDOS = Path(carpeta_temporal) / "pedidos.xlsx"
    with patch.object(datos, "_ARCHIVO_PEDIDOS", RUTA_PEDIDOS):
        if RUTA_PEDIDOS.exists():
            RUTA_PEDIDOS.unlink()

        hilos = [
            threading.Thread(target=registrar_pedido_de_cliente, args=(numero,))
            for numero in range(NUM_CLIENTES)
        ]
        for hilo in hilos:
            hilo.start()
        for hilo in hilos:
            hilo.join()

        if errores:
            raise AssertionError(f"Falló el registro de pedidos: {errores}")

        cantidad_codigos_distintos = len(set(codigos))
        print(
            f"Códigos distintos: {cantidad_codigos_distintos} "
            f"de {NUM_CLIENTES} clientes."
        )
        assert cantidad_codigos_distintos == NUM_CLIENTES

        libro = load_workbook(RUTA_PEDIDOS, read_only=True, data_only=True)
        try:
            hoja = libro.active
            encabezados = tuple(next(hoja.iter_rows(values_only=True)))
            filas = list(hoja.iter_rows(min_row=2, values_only=True))
        finally:
            libro.close()

        assert encabezados == datos._ENCABEZADOS
        assert len(filas) == NUM_CLIENTES
        assert {fila[0] for fila in filas} == set(codigos)
        assert {fila[3] for fila in filas} == {
            f"cliente-{numero}" for numero in range(NUM_CLIENTES)
        }

print("Prueba correcta: los 10 pedidos concurrentes quedaron guardados.")
