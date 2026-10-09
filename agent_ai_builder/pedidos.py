import json
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from .datos import Producto


_ARCHIVO_PEDIDOS = (
    Path(__file__).resolve().parent.parent / "datos" / "pedidos.xlsx"
)
_ENCABEZADOS = (
    "ID",
    "Fecha y hora",
    "Productos",
    "Total referencial (RD$)",
    "Estado",
)


def registrar_pedido(
    items: list[dict[str, str | int]],
    productos: dict[str, Producto],
) -> str:
    detalles = []
    total = 0

    for item in items:
        nombre_producto = item["producto"].strip()
        producto = productos.get(nombre_producto.lower())
        if producto is None:
            return f"El producto '{nombre_producto}' no está en el menú."

        cantidad = item["cantidad"]
        if type(cantidad) is not int or cantidad <= 0:
            return "La cantidad de cada producto debe ser un entero mayor que cero."

        precio = producto["price"]
        subtotal = precio * cantidad
        total += subtotal
        detalles.append(
            {
                "producto": producto["name"],
                "cantidad": cantidad,
                "precio_unitario": precio,
                "subtotal": subtotal,
            }
        )

    if not detalles:
        return "No se puede registrar un pedido sin productos."

    pedido_id = str(uuid4())
    _ARCHIVO_PEDIDOS.parent.mkdir(parents=True, exist_ok=True)
    if _ARCHIVO_PEDIDOS.exists():
        libro = load_workbook(_ARCHIVO_PEDIDOS)
        hoja = libro.active
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "pedidos"
        hoja.append(_ENCABEZADOS)

    try:
        hoja.append(
            (
                pedido_id,
                datetime.now().astimezone().isoformat(timespec="seconds"),
                json.dumps(detalles, ensure_ascii=False),
                total,
                "Pendiente de confirmación",
            )
        )
        libro.save(_ARCHIVO_PEDIDOS)
    finally:
        libro.close()

    return (
        f"Pedido {pedido_id} registrado con total referencial de RD${total}. "
        "El negocio debe confirmar disponibilidad y precio."
    )
