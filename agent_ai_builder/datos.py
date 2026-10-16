import json
import threading
from datetime import datetime
from pathlib import Path
from typing import TypedDict
from uuid import uuid4

from openpyxl import Workbook, load_workbook


_ARCHIVO_DATOS = (
    Path(__file__).resolve().parent.parent
    / "datos"
    / "datos_negocio_cafe_santo_domingo.xlsx"
)
_NOMBRE_NEGOCIO = "Café Santo Domingo"
_ARCHIVO_PEDIDOS = (
    Path(__file__).resolve().parent.parent / "datos" / "pedidos.xlsx"
)
_ENCABEZADOS = (
    "ID",
    "Fecha y hora",
    "Cliente",
    "Contacto",
    "Detalle",
    "Total referencial (RD$)",
    "Estado",
)
_ENCABEZADOS_ANTERIORES = (
    "ID",
    "Fecha y hora",
    "Productos",
    "Total referencial (RD$)",
    "Estado",
)
_candado_pedidos = threading.Lock()


class Producto(TypedDict):
    name: str
    price: int | float


class Categoria(TypedDict):
    category: str
    items: list[Producto]


class Informacion(TypedDict):
    topic: str
    details: str


class DatosNegocio(TypedDict):
    business_name: str
    catalog: list[Categoria]
    information: list[Informacion]


def cargar_datos_negocio() -> DatosNegocio:
    """Carga el catálogo y la información del negocio desde la hoja de cálculo."""
    libro = load_workbook(_ARCHIVO_DATOS, read_only=True, data_only=True)
    try:
        productos_por_categoria: dict[str, list[Producto]] = {}
        for producto, precio, categoria in libro["productos"].iter_rows(
            min_row=2,
            values_only=True,
        ):
            productos_por_categoria.setdefault(categoria, []).append(
                {"name": producto, "price": precio}
            )

        informacion = [
            {"topic": tema, "details": detalle}
            for tema, detalle in libro["informacion"].iter_rows(
                min_row=2,
                values_only=True,
            )
        ]
    finally:
        libro.close()

    catalogo = [
        {"category": categoria, "items": productos}
        for categoria, productos in productos_por_categoria.items()
    ]
    return {
        "business_name": _NOMBRE_NEGOCIO,
        "catalog": catalogo,
        "information": informacion,
    }


DATOS_NEGOCIO = cargar_datos_negocio()


def registrar_pedido(
    items: list[dict[str, str | int]],
    productos: dict[str, Producto],
    contacto: str = "",
) -> str:
    with _candado_pedidos:
        detalles = []
        total = 0

        for item in items:
            nombre_producto = item["producto"].strip()
            producto = productos.get(nombre_producto.lower())
            if producto is None:
                return f"El producto '{nombre_producto}' no está en el menú."

            cantidad = item["cantidad"]
            if type(cantidad) is not int or cantidad <= 0:
                return (
                    "La cantidad de cada producto debe ser "
                    "un entero mayor que cero."
                )

            precio = producto["price"]
            subtotal = precio * cantidad
            total += subtotal
            detalles.append(
                {
                    "producto": producto["name"],
                    "cantidad": cantidad,
                    "precio_unitario": precio,
                    "subtotal": subtotal,
                }
            )

        if not detalles:
            return "No se puede registrar un pedido sin productos."

        pedido_id = str(uuid4())
        _ARCHIVO_PEDIDOS.parent.mkdir(parents=True, exist_ok=True)
        if _ARCHIVO_PEDIDOS.exists():
            libro = load_workbook(_ARCHIVO_PEDIDOS)
            hoja = libro.active
            encabezados = tuple(
                hoja.cell(row=1, column=columna).value
                for columna in range(1, hoja.max_column + 1)
            )
            if encabezados == _ENCABEZADOS_ANTERIORES:
                hoja.insert_cols(3, amount=2)
                for columna, encabezado in enumerate(_ENCABEZADOS, start=1):
                    hoja.cell(row=1, column=columna, value=encabezado)
            elif encabezados != _ENCABEZADOS:
                libro.close()
                raise ValueError(
                    "La hoja de pedidos tiene encabezados incompatibles; "
                    "no se puede registrar el pedido con seguridad."
                )
        else:
            libro = Workbook()
            hoja = libro.active
            hoja.title = "pedidos"
            hoja.append(_ENCABEZADOS)

        try:
            hoja.append(
                (
                    pedido_id,
                    datetime.now().astimezone().isoformat(timespec="seconds"),
                    "",
                    contacto,
                    json.dumps(detalles, ensure_ascii=False),
                    total,
                    "Confirmado por el cliente",
                )
            )
            libro.save(_ARCHIVO_PEDIDOS)
        finally:
            libro.close()

        return (
            f"Pedido {pedido_id} confirmado por el cliente y registrado con "
            f"total referencial de RD${total}. El negocio debe confirmar "
            "disponibilidad y precio."
        )
